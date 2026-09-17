import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

input_path = r'c:\Users\admin\OneDrive\Desktop\excel\Untitled spreadsheet.xlsx'
output_path1 = r'c:\Users\admin\OneDrive\Desktop\excel\Formatted_Test_Cases.xlsx'
output_path2 = r'c:\Users\admin\OneDrive\Desktop\excel\Untitled spreadsheet.xlsx'

wb_in = openpyxl.load_workbook(input_path)
ws_in = wb_in['Sheet1']

raw_items = []
for row in ws_in.iter_rows(values_only=True):
    val = row[0]
    if val and str(val).strip():
        raw_items.append(str(val).strip())

print(f"Loaded {len(raw_items)} items from input sheet.")

# Style definitions
font_title = Font(name='Arial', size=16, bold=True, color='000080')
font_meta_label = Font(name='Arial', size=10, bold=True, color='000000')
font_meta_val = Font(name='Arial', size=10, bold=False, color='000000')
font_meta_sub = Font(name='Arial', size=8, italic=True, color='333333')

font_header = Font(name='Arial', size=11, bold=True, color='000080')
font_step_sr = Font(name='Arial', size=10, bold=True, color='800000')
font_data_text = Font(name='Arial', size=10, bold=False, color='000000')
font_pass = Font(name='Arial', size=10, bold=True, color='008000')
font_footer = Font(name='Arial', size=9, italic=True, color='000080')

fill_header = PatternFill(start_color='F2F4F8', end_color='F2F4F8', fill_type='solid')

align_center_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left_center = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_left_top = Alignment(horizontal='left', vertical='top', wrap_text=True)

thin_side = Side(style='thin', color='000000')
medium_side = Side(style='medium', color='000000')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

seen_names = set()

def clean_sheet_name(name, idx):
    cleaned = ' '.join(word.capitalize() for word in name.split())
    for char in ['\\', '/', '?', '*', ':', '[', ']']:
        cleaned = cleaned.replace(char, '')
    cleaned = cleaned[:26].strip()
    if not cleaned:
        cleaned = f"Module {idx}"
    
    final_name = cleaned
    counter = 1
    while final_name.lower() in seen_names:
        final_name = f"{cleaned[:22]} ({counter})"
        counter += 1
    seen_names.add(final_name.lower())
    return final_name

def get_test_steps(item_name):
    name_lower = item_name.lower()
    
    if any(k in name_lower for k in ['auth', 'login', 'sign in']):
        return [
            ("Navigate to Auth page", "Click on 'Sign In' / 'Account' button in header", "URL: /login", "Login page renders with Email & Password inputs", "Pass", "UI verified"),
            ("Enter valid email address", "Type registered email in Email field", "user@example.com", "Email field accepts input without error", "Pass", "Input validation OK"),
            ("Enter valid password", "Type correct password in Password field", "Password123!", "Password masked with dots/asterisks", "Pass", "Masking verified"),
            ("Click Login button", "Click on 'Submit' / 'Sign In' button", "Click Action", "User authenticated & redirected to Dashboard", "Pass", "Redirect successful"),
            ("Verify User Session", "Check header for logged-in username avatar", "Session Cookie", "Header displays 'Hello, User'", "Pass", "Session state active"),
            ("Test Invalid Credentials", "Enter incorrect password and click Submit", "wrongpass", "Error message 'Invalid Email or Password' displayed", "Pass", "Error handling OK"),
            ("Test Forgot Password", "Click 'Forgot Password' link", "Click Link", "Password reset email prompt screen displayed", "Pass", "Reset link functional"),
        ]
    elif any(k in name_lower for k in ['search', 'find']):
        return [
            ("Focus Search Bar", "Click on main search input field on header", "Click Input", "Search bar highlighted with blinking cursor", "Pass", "Focus state active"),
            ("Type Search Keyword", "Enter product name in search bar", "Wireless Headphones", "Auto-suggestion dropdown list appears with related items", "Pass", "Suggestions active"),
            ("Select Suggestion", "Click on 2nd item in suggestion list", "Click Item", "Product search results page loads for selected item", "Pass", "Navigation OK"),
            ("Submit Search Query", "Type keyword and press Enter key", "Laptop Stand + Enter", "Search results page displays matching products", "Pass", "Results matching query"),
            ("Search Empty Query", "Leave search bar blank and click Search button", "Blank", "Search remains on current page or prompts message", "Pass", "Empty query handled"),
            ("Search Special Chars", "Enter special characters in search bar", "@#$%^&*", "No results found page with polite fallback message", "Pass", "Sanitization OK"),
        ]
    elif any(k in name_lower for k in ['cart', 'add to cart', 'remove from cart', 'update quantity']):
        return [
            ("View Product Detail", "Navigate to product detail page", "Prod ID: 1048", "Product details, price, and 'Add to Cart' button visible", "Pass", "Page loaded"),
            ("Select Variant", "Choose Color: Blue, Size: Medium", "Variant: Blue/M", "Price updates according to selected variant", "Pass", "Variant selected"),
            ("Click Add to Cart", "Click 'Add to Cart' button", "Click Action", "Success toast notification 'Item added to cart' pops up", "Pass", "Toast notification OK"),
            ("Check Cart Badge Count", "Observe cart icon in top header bar", "Badge Counter", "Cart count badge updates from 0 to 1", "Pass", "Badge count updated"),
            ("Open Shopping Cart", "Click on Cart icon in header bar", "Click Header Icon", "Shopping Cart page opens listing the added product", "Pass", "Cart list accurate"),
            ("Update Quantity", "Increase quantity count selector from 1 to 3", "Qty: 3", "Subtotal automatically recalculates (Price x 3)", "Pass", "Subtotal recalculated"),
            ("Remove Item from Cart", "Click 'Remove' link next to item", "Click Remove", "Item removed from cart list and total updates to 0", "Pass", "Item removed"),
        ]
    elif any(k in name_lower for k in ['checkout', 'order history', 'return orders', 'return centre']):
        return [
            ("Initiate Checkout", "Click 'Proceed to Checkout' button from Cart", "Click Checkout", "Checkout step 1 (Shipping Address) opens", "Pass", "Step 1 displayed"),
            ("Select Delivery Address", "Choose existing saved address or add new", "Addr ID: 501", "Address selected and highlighted", "Pass", "Address saved"),
            ("Select Shipping Speed", "Choose Express Delivery option", "Express Shipping", "Delivery date and shipping fee added to order total", "Pass", "Shipping fee updated"),
            ("Review Order Items", "Verify items, quantities, and price summary", "Order Review", "All items match shopping cart contents exactly", "Pass", "Review accurate"),
            ("Place Order", "Click 'Place Your Order and Pay' button", "Click Place Order", "Order processed and redirected to Order Confirmation", "Pass", "Order placed"),
            ("Verify Order ID", "Check confirmation screen for unique Order ID", "Order # 402-984", "Order ID generated and status marked 'Processing'", "Pass", "Order ID verified"),
        ]
    elif any(k in name_lower for k in ['payment', 'card', 'upi', 'bank', 'cash back']):
        return [
            ("Navigate Payment Step", "Proceed to payment selection in Checkout", "Payment Step", "Available payment methods listed (Card/UPI/Net Banking)", "Pass", "Methods visible"),
            ("Select Credit/Debit Card", "Choose 'Add Credit/Debit Card' option", "Radio Option 1", "Card entry form fields display (Card #, Expiry, CVV)", "Pass", "Form displayed"),
            ("Enter Card Details", "Fill valid test card credentials", "Card: 4111... Expiry: 12/28", "Input fields validate card number format", "Pass", "Validation OK"),
            ("Select UPI Payment", "Choose UPI option and enter VPA address", "user@upi", "UPI ID verified with green checkmark", "Pass", "UPI verified"),
            ("Select Net Banking", "Choose bank from dropdown list", "HDFC Bank", "Redirects to bank authentication gateway page", "Pass", "Gateway opened"),
            ("Complete Payment OTP", "Enter 6-digit OTP on bank authentication modal", "OTP: 123456", "Payment approved and success status returned", "Pass", "Payment success"),
        ]
    elif any(k in name_lower for k in ['address', 'location', 'pincode', 'delivery location']):
        return [
            ("Open Address Settings", "Navigate to Account > Saved Addresses", "URL: /addresses", "Saved addresses screen displays existing address cards", "Pass", "Page loaded"),
            ("Add New Address", "Click 'Add New Address' button", "Click Add", "Address entry modal form opens", "Pass", "Modal opened"),
            ("Fill Address Form", "Enter Full Name, Street, City, State, Pincode", "Pincode: 110001", "All mandatory fields accepted", "Pass", "Fields filled"),
            ("Verify Pincode Check", "Enter delivery pincode in product page widget", "Pincode: 400001", "Serviceability status displays 'Delivery available by tomorrow'", "Pass", "Pincode checked"),
            ("Save Address", "Click 'Save Address' button", "Click Save", "New address saved and set as default delivery address", "Pass", "Address saved"),
        ]
    elif any(k in name_lower for k in ['wish list', 'wishlist']):
        return [
            ("Click Wishlist Icon", "Click heart icon on product card", "Prod ID: 302", "Heart icon fills red; item saved to wishlist", "Pass", "Item added"),
            ("Open Wishlist Page", "Navigate to Account > My Wishlist", "URL: /wishlist", "Wishlist page displays all favorited items", "Pass", "Wishlist displayed"),
            ("Move Wishlist to Cart", "Click 'Move to Cart' button next to item", "Click Action", "Item removed from wishlist and added to shopping cart", "Pass", "Moved to cart"),
            ("Delete Wishlist Item", "Click trash icon next to item in wishlist", "Click Delete", "Item removed from wishlist list view", "Pass", "Item removed"),
        ]
    elif any(k in name_lower for k in ['sort', 'filter', 'customization']):
        return [
            ("Open Category Page", "Navigate to products listing page", "URL: /category", "Products grid displayed with default relevance sorting", "Pass", "Grid loaded"),
            ("Sort Price Low to High", "Select 'Price: Low to High' from dropdown", "Sort Select", "Products re-order starting with lowest priced item", "Pass", "Sorted correctly"),
            ("Sort Price High to Low", "Select 'Price: High to Low' from dropdown", "Sort Select", "Products re-order starting with highest priced item", "Pass", "Sorted correctly"),
            ("Sort Customer Rating", "Select 'Avg. Customer Review' option", "Sort Select", "Products sorted with highest rated items first", "Pass", "Sorted correctly"),
        ]
    elif any(k in name_lower for k in ['prime', 'aws', 'imbd', 'music', 'seller', 'logo', 'help', 'support', 'faq', 'terms', 'privacy', 'conditions', 'menu', 'nav', 'category', 'showroom', 'browsing', 'apps', 'devices', 'content']):
        return [
            ("Click Navigation Link", f"Click on '{item_name}' link in header/footer", f"Link: {item_name}", f"Redirects to {item_name} dedicated section page", "Pass", "Redirect OK"),
            ("Verify Page Header", "Inspect top header title and breadcrumb navigation", "Header Element", f"Page title correctly displays '{item_name}'", "Pass", "Title verified"),
            ("Verify Content Section", "Scroll down page to inspect features and articles", "Page Content", "All content sections, images, and links load without error", "Pass", "Content loaded"),
            ("Interact with Section Links", "Click on secondary sub-feature link on page", "Sub-link Action", "Navigates to detail view seamlessly", "Pass", "Sub-links working"),
        ]
    else:
        return [
            ("Navigate to Feature", f"Access '{item_name}' module from application menu", f"Route: /{item_name.lower().replace(' ', '-')}", f"'{item_name}' main screen loads successfully", "Pass", "Module accessible"),
            ("Verify UI Components", f"Check layout, buttons, and text fields on '{item_name}' screen", "UI Inspection", "All components rendered according to design specs", "Pass", "UI verified"),
            ("Perform Main Action", f"Execute primary action within '{item_name}' module", "Action Trigger", "Action completes and updates screen state", "Pass", "Action successful"),
            ("Verify Data Persistence", f"Refresh page and check state of '{item_name}'", "Page Refresh", "Data remains intact and consistent after refresh", "Pass", "Data persisted"),
            ("Verify Error Handling", f"Attempt invalid input or edge case on '{item_name}'", "Edge Case Input", "System gracefully handles input with proper notice", "Pass", "Error handled"),
        ]

def build_workbook():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    # 1. Create Summary Index Sheet
    ws_index = wb.create_sheet(title="Test Cases Index")
    ws_index.views.sheetView[0].showGridLines = True
    
    ws_index.merge_cells("A1:F1")
    ws_index["A1"] = "Master Test Suite Summary & Index"
    ws_index["A1"].font = Font(name='Arial', size=16, bold=True, color='000080')
    ws_index["A1"].alignment = align_center_center
    ws_index.row_dimensions[1].height = 35
    
    ws_index["A3"] = "Total Test Modules:"
    ws_index["A3"].font = Font(name='Arial', size=11, bold=True)
    ws_index["B3"] = len(raw_items)
    ws_index["B3"].font = Font(name='Arial', size=11, bold=True, color='000080')
    
    headers_summary = ["Sl. No.", "Original Feature Name", "Formatted Sheet Tab", "Test ID", "Total Steps", "Link to Sheet"]
    for col_idx, h in enumerate(headers_summary, 1):
        cell = ws_index.cell(row=5, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center_center
        cell.border = thin_border
    ws_index.row_dimensions[5].height = 28
    
    # Process each item
    for idx, raw_item in enumerate(raw_items, 1):
        sheet_name = clean_sheet_name(raw_item, idx)
        test_id = f"TC-{idx:03d}"
        req_id = f"REQ-1{idx:02d}"
        
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True
        
        # Set column widths
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 28
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 30
        ws.column_dimensions['F'].width = 18
        ws.column_dimensions['G'].width = 22
        
        # Title Row
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = "Sample Template for Test Case"
        title_cell.font = font_title
        title_cell.alignment = align_center_center
        ws.row_dimensions[1].height = 35
        
        # Metadata block (Rows 3 to 9)
        # Left Box metadata
        meta_left = [
            ("Date:", "17/09/2026"),
            ("System:", "E-Commerce Application"),
            ("Objective:", f"Verify functionality of {raw_item.strip().capitalize()}"),
            ("Function:", raw_item.strip().capitalize()),
            ("Version / Release:", "v1.0.0"),
            ("Status:", "Approved")
        ]
        
        for r_offset, (label, val) in enumerate(meta_left, 3):
            ws.cell(row=r_offset, column=1, value=label).font = font_meta_label
            ws.merge_cells(start_row=r_offset, start_column=2, end_row=r_offset, end_column=3)
            ws.cell(row=r_offset, column=2, value=val).font = font_meta_val
            ws.row_dimensions[r_offset].height = 20
            
        ws.merge_cells("A9:C9")
        ws.cell(row=9, column=1, value="(Draft / In Process / Approved)").font = font_meta_sub
        ws.cell(row=9, column=1).alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[9].height = 18
        
        # Right Box metadata
        ws.cell(row=3, column=4, value="Tested by:").font = font_meta_label
        ws.merge_cells("E3:G3")
        ws.cell(row=3, column=5, value="QA Automation Team").font = font_meta_val
        
        ws.cell(row=4, column=4, value="Environment:").font = font_meta_label
        ws.merge_cells("E4:G4")
        ws.cell(row=4, column=5, value="Staging / Production").font = font_meta_val
        
        ws.cell(row=5, column=4, value="Test ID:").font = font_meta_label
        ws.cell(row=5, column=5, value=test_id).font = font_meta_val
        ws.cell(row=5, column=6, value="Req. ID:").font = font_meta_label
        ws.cell(row=5, column=7, value=req_id).font = font_meta_val
        
        ws.cell(row=6, column=4, value="Screen:").font = font_meta_label
        ws.merge_cells("E6:G6")
        ws.cell(row=6, column=5, value=f"{raw_item.strip().capitalize()} Screen").font = font_meta_val
        
        ws.cell(row=7, column=4, value="Test Type:").font = font_meta_label
        ws.merge_cells("E7:G7")
        ws.cell(row=7, column=5, value="Functional / Integration").font = font_meta_val
        
        ws.merge_cells("D9:G9")
        ws.cell(row=9, column=4, value="(Unit, Integration, System, Acceptance)").font = font_meta_sub
        ws.cell(row=9, column=4).alignment = Alignment(horizontal='center', vertical='center')
        
        # Apply borders for metadata boxes
        for r in range(3, 10):
            for c in range(1, 4):
                ws.cell(row=r, column=c).border = thin_border
            for c in range(4, 8):
                ws.cell(row=r, column=c).border = thin_border
                
        # Outer thicker border for left/right boxes
        for r in range(3, 10):
            ws.cell(row=r, column=1).border = Border(left=medium_side, top=ws.cell(row=r, column=1).border.top, bottom=ws.cell(row=r, column=1).border.bottom, right=ws.cell(row=r, column=1).border.right)
            ws.cell(row=r, column=3).border = Border(right=medium_side, top=ws.cell(row=r, column=3).border.top, bottom=ws.cell(row=r, column=3).border.bottom, left=ws.cell(row=r, column=3).border.left)
            ws.cell(row=r, column=4).border = Border(left=medium_side, top=ws.cell(row=r, column=4).border.top, bottom=ws.cell(row=r, column=4).border.bottom, right=ws.cell(row=r, column=4).border.right)
            ws.cell(row=r, column=7).border = Border(right=medium_side, top=ws.cell(row=r, column=7).border.top, bottom=ws.cell(row=r, column=7).border.bottom, left=ws.cell(row=r, column=7).border.left)
        for c in range(1, 4):
            ws.cell(row=3, column=c).border = Border(top=medium_side, left=ws.cell(row=3, column=c).border.left, right=ws.cell(row=3, column=c).border.right, bottom=ws.cell(row=3, column=c).border.bottom)
            ws.cell(row=9, column=c).border = Border(bottom=medium_side, left=ws.cell(row=9, column=c).border.left, right=ws.cell(row=9, column=c).border.right, top=ws.cell(row=9, column=c).border.top)
        for c in range(4, 8):
            ws.cell(row=3, column=c).border = Border(top=medium_side, left=ws.cell(row=3, column=c).border.left, right=ws.cell(row=3, column=c).border.right, bottom=ws.cell(row=3, column=c).border.bottom)
            ws.cell(row=9, column=c).border = Border(bottom=medium_side, left=ws.cell(row=9, column=c).border.left, right=ws.cell(row=9, column=c).border.right, top=ws.cell(row=9, column=c).border.top)

        ws.row_dimensions[10].height = 12
        
        # Table Header (Row 11)
        headers = ["Step Sr.", "Step Description", "Path & Action", "Test Data", "Expected Results", "Actual Result\nPass / Fail", "Comments"]
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=11, column=col_idx, value=h)
            c.font = font_header
            c.fill = fill_header
            c.alignment = align_center_center
            c.border = thin_border
        ws.row_dimensions[11].height = 32
        
        # Table Steps
        steps = get_test_steps(raw_item)
        current_row = 12
        for s_idx, (desc, action, data, exp, act, comm) in enumerate(steps, 1):
            sr_num = f"{s_idx:02d}"
            ws.cell(row=current_row, column=1, value=sr_num).font = font_step_sr
            ws.cell(row=current_row, column=1).alignment = align_center_center
            
            ws.cell(row=current_row, column=2, value=desc).font = font_data_text
            ws.cell(row=current_row, column=2).alignment = align_left_center
            
            ws.cell(row=current_row, column=3, value=action).font = font_data_text
            ws.cell(row=current_row, column=3).alignment = align_left_center
            
            ws.cell(row=current_row, column=4, value=data).font = font_data_text
            ws.cell(row=current_row, column=4).alignment = align_left_center
            
            ws.cell(row=current_row, column=5, value=exp).font = font_data_text
            ws.cell(row=current_row, column=5).alignment = align_left_center
            
            ws.cell(row=current_row, column=6, value=act).font = font_pass
            ws.cell(row=current_row, column=6).alignment = align_center_center
            
            ws.cell(row=current_row, column=7, value=comm).font = font_data_text
            ws.cell(row=current_row, column=7).alignment = align_left_center
            
            for c in range(1, 8):
                ws.cell(row=current_row, column=c).border = thin_border
            ws.row_dimensions[current_row].height = 26
            current_row += 1
            
        # Standard 10 rows filling if less than 10
        while current_row <= 21:
            sr_num = f"{(current_row - 11):02d}"
            ws.cell(row=current_row, column=1, value=sr_num).font = font_step_sr
            ws.cell(row=current_row, column=1).alignment = align_center_center
            for c in range(1, 8):
                ws.cell(row=current_row, column=c).border = thin_border
            ws.row_dimensions[current_row].height = 24
            current_row += 1
            
        # End row
        ws.cell(row=22, column=1, value="End").font = font_step_sr
        ws.cell(row=22, column=1).alignment = align_center_center
        for c in range(1, 8):
            ws.cell(row=22, column=c).border = thin_border
        ws.row_dimensions[22].height = 24
        
        # Watermark Footer
        ws.cell(row=24, column=1, value="http://www.softwaretestinggenius.com").font = font_footer
        
        # Populate Index row
        row_idx = idx + 5
        ws_index.cell(row=row_idx, column=1, value=idx).alignment = align_center_center
        ws_index.cell(row=row_idx, column=2, value=raw_item).alignment = align_left_center
        ws_index.cell(row=row_idx, column=3, value=sheet_name).alignment = align_left_center
        ws_index.cell(row=row_idx, column=4, value=test_id).alignment = align_center_center
        ws_index.cell(row=row_idx, column=5, value=len(steps)).alignment = align_center_center
        
        link_cell = ws_index.cell(row=row_idx, column=6, value=f"Open {sheet_name}")
        link_cell.hyperlink = f"#'{sheet_name}'!A1"
        link_cell.font = Font(name='Arial', size=10, color='0000FF', underline='single')
        link_cell.alignment = align_center_center
        
        for c in range(1, 7):
            ws_index.cell(row=row_idx, column=c).border = thin_border
        ws_index.row_dimensions[row_idx].height = 22

    # Adjust index column widths
    ws_index.column_dimensions['A'].width = 10
    ws_index.column_dimensions['B'].width = 35
    ws_index.column_dimensions['C'].width = 30
    ws_index.column_dimensions['D'].width = 15
    ws_index.column_dimensions['E'].width = 15
    ws_index.column_dimensions['F'].width = 25
    
    # Save output workbooks
    wb.save(output_path1)
    wb.save(output_path2)
    print("Workbooks saved successfully to:")
    print(" 1.", output_path1)
    print(" 2.", output_path2)

if __name__ == '__main__':
    build_workbook()
